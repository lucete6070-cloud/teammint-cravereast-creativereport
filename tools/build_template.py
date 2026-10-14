"""공식 양식에 'TOP10 상세표' 슬라이드를 추가해 template/EAST_template.pptx 를 만든다.

- 입력: template/EAST_template_official.pptx (공식 양식 원본), reference/ 의 9월 FN pptx (표 서식 출처)
- FN 의 표 슬라이드(2개 표)를 복제하되, 표 안의 값은 모두 비운다 → template/ 은 git 에 올려도 광고주 데이터가 없다.
- 한 번만 돌리면 된다 (양식이 바뀔 때만 다시).
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from pptx import Presentation  # noqa: E402
from pptx_util import new_slide, clone_shapes, set_para_text  # noqa: E402

FN = ROOT / "reference" / "2026-09_EAST-2_FN.pptx"
OFFICIAL = ROOT / "template" / "EAST_template_official.pptx"
OUT = ROOT / "template" / "EAST_template.pptx"
TABLE_SLIDE_INDEX = 6  # FN 7번째 슬라이드: MY OG 상세표 (표 2개)


def main():
    fn = Presentation(str(FN))
    src = fn.slides[TABLE_SLIDE_INDEX]
    prs = Presentation(str(OFFICIAL))
    dst = new_slide(prs, "MY - OG")
    clone_shapes(src, dst)
    for sh in dst.shapes:
        if sh.has_table:
            for ri, row in enumerate(sh.table.rows):
                if ri == 0:
                    continue  # 헤더는 유지
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        set_para_text(p, "")
                    # 우수/저조 하이라이트는 생성 시 다시 칠하므로 배경을 지운다
                    tcPr = cell._tc.tcPr
                    if tcPr is not None:
                        for el in list(tcPr):
                            if el.tag.endswith("}solidFill"):
                                tcPr.remove(el)
        elif sh.has_text_frame:
            t = sh.text_frame.text
            if t.startswith("SKIN1004"):
                set_para_text(sh.text_frame.paragraphs[0], "SKIN1004 (국가) (상시/프로모션) 캠페인별 우수 소재 상세 효율")
            elif t.startswith("▼"):
                set_para_text(sh.text_frame.paragraphs[0], "▼ <캠페인디테일> 소재 지출 금액 TOP 10")
    prs.save(str(OUT))
    print("만들었습니다:", OUT)


if __name__ == "__main__":
    main()
