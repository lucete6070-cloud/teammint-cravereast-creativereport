"""[테스트용] 9월 FN(reference/)에서 소재 이미지를 뽑아 images/소재명.png|jpg 로 저장한다.

images/ 가 비어 있을 때 생성기를 시험하려고 쓴다. 실제 월간 실행에서는 보조가
광고관리자에서 받은 이미지를 images/ 에 넣는다.
"""
import io
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx_read import _walk, slide_cards  # noqa: E402
from sheet import read_csv, parse_blocks  # noqa: E402

FN = ROOT / "reference" / "2026-09_EAST-2_FN.pptx"
OUT = ROOT / "images"


def csv_names():
    """input/ CSV 의 소재명 전체 (FN 에서 잘려 있는 이름을 보정하는 데 사용)."""
    names = set()
    for p in (ROOT / "input").glob("*.csv"):
        for b in parse_blocks(read_csv(p)):
            names.update(r.name for r in b.rows)
    return names


def fix_name(name, known):
    low = {k.casefold(): k for k in known}
    if name.casefold() in low:
        return low[name.casefold()]
    cand = [k for k in known if k.casefold().startswith(name.casefold())]
    return cand[0] if len(cand) == 1 else name


def iter_fn_images(prs, known=()):
    """FN 슬라이드의 (슬라이드번호, 소재명, PIL 이미지) — 크롭 반영. compare_images.py 도 이 함수를 쓴다."""
    for n, slide in enumerate(prs.slides, 1):
        lay = slide.slide_layout.name
        if not (lay.endswith(" - OG") or lay.endswith(" - PR")):
            continue
        cards = slide_cards(slide)
        if not cards:
            continue
        pics = [(sh, l, t, w, h) for sh, l, t, w, h in _walk(slide.shapes)
                if sh.shape_type == 13 and w > 1.5 and abs(w - h) < 0.8]
        for sh, l, t, w, h in pics:
            cx = l + w / 2
            card = min(cards, key=lambda c: abs(c["x"] - cx))
            if abs(card["x"] - cx) > 1.3:
                yield n, None, None
                continue
            name = fix_name(re.sub(r"[\s\x0b]+", "", card["name"]), known)
            im = Image.open(io.BytesIO(sh.image.blob))
            cl, cr, ct, cb = sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom
            if any((cl, cr, ct, cb)):
                W, H = im.size
                im = im.crop((int(W * cl), int(H * ct), int(W * (1 - cr)), int(H * (1 - cb))))
            yield n, name, im


def main():
    known = csv_names()
    OUT.mkdir(exist_ok=True)
    prs = Presentation(str(FN))
    saved = 0
    for n, name, im in iter_fn_images(prs, known):
        if not name:
            print("건너뜀:", n)
            continue
        if im.format == "JPEG" or name.endswith(".jpg"):
            im.convert("RGB").save(OUT / f"{name}.jpg", "JPEG", quality=92)
        else:
            im.save(OUT / f"{name}.png", "PNG")
        saved += 1
    print(f"이미지 {saved}장 저장 → {OUT}")


if __name__ == "__main__":
    main()
