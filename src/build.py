"""양식(template/)을 복제·채워서 국가×유형별 pptx 를 만든다.

- 코멘트 박스·콜아웃은 비워 둔다 (코멘트는 별도 스킬 담당).
- 콘텐츠 슬라이드 노트에는 코멘트 근거(슬롯·시트 지표·전월 기록)를 적는다 (notes.py).
- 이미지는 양식의 '소재 이미지' 사각형 위치·크기에 맞춰 정사각 크롭 후 교체. 없으면 사각형을 남기고 기록한다.
- 텍스트는 run 단위로 교체해 서식을 보존한다.
"""
import copy
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches

import notes
from pptx_util import (A_NS, clear_text, clone_shapes, delete_slide, move_slide_to_end_order, new_slide,
                       set_lines, set_para_text, split_sku)

IMG_EXT = (".png", ".jpg", ".jpeg")
GREEN, PINK = "E2F0D9", "FADADD"   # 표 하이라이트: 우수=연두, 저조=분홍 [추정: SPEC 3장 기준, FN 은 우수만 칠함]
SLOT_TEMPLATE_INDEX = {3: 1, 4: 2, 5: 3}   # 양식 슬라이드 번호(0부터): 3/4/5슬롯
TABLE_TEMPLATE_INDEX = 17
EOD_INDEX = 16


class ImageIndex:
    """images/ 폴더에서 파일명(확장자 뺀 것) = 소재명 으로 찾는다. 대소문자는 구분하지 않는다.
    images_manual/ (images/ 와 같은 위치)에 같은 이름의 파일이 있으면 그것이 우선이고, 자동 수집이 덮어쓰지 않는다."""

    def __init__(self, folder):
        self.map = {}
        self.manual_keys = set()
        folder = Path(folder)
        if folder.exists():
            for p in folder.iterdir():
                if p.suffix.lower() in IMG_EXT:
                    self.map[p.stem.casefold()] = p
        manual = folder.parent / "images_manual" if folder.name == "images" else None
        if manual is not None and manual.exists():
            for p in manual.iterdir():
                if p.suffix.lower() in IMG_EXT:
                    self.map[p.stem.casefold()] = p
                    self.manual_keys.add(p.stem.casefold())

    def find(self, name):
        return self.map.get(name.casefold())


# ----------------------------------------------------------------- 도형 찾기
def _cx(s):
    return s.left + s.width / 2


def find_slots(slide):
    top = list(slide.shapes)
    rects = sorted([s for s in top if s.has_text_frame and s.text_frame.text.strip() == "소재 이미지"], key=lambda s: s.left)
    labels = [s for s in top if s.has_text_frame and s.text_frame.text.strip() == "캠페인명"]
    groups = [s for s in top if s.shape_type == 6 and any(
        c.has_text_frame and c.text_frame.text.strip().startswith("지출") for c in s.shapes)]
    slots = []
    for r in rects:
        lab = min(labels, key=lambda l: abs(_cx(l) - _cx(r)))
        grp = min(groups, key=lambda g: abs(_cx(g) - _cx(r)))
        kpi = next(c for c in grp.shapes if c.has_text_frame and c.text_frame.text.strip().startswith("지출"))
        nm = next(c for c in grp.shapes if c.has_text_frame and c.text_frame.text.strip().startswith("소재명"))
        slots.append(dict(rect=r, label=lab, group=grp, kpi=kpi, name=nm))
    return slots


def _remove(sh):
    sh._element.getparent().remove(sh._element)


def put_image(slide, rect, path):
    """rect 자리에 정사각 크롭한 그림을 넣고 rect 를 지운다. 그림은 rect 의 겹침 순서를 이어받는다."""
    pic = slide.shapes.add_picture(str(path), rect.left, rect.top, rect.width, rect.height)
    with Image.open(path) as im:
        w, h = im.size
    if w > h:
        pic.crop_left = pic.crop_right = (w - h) / w / 2
    elif h > w:
        pic.crop_top = pic.crop_bottom = (h - w) / h / 2
    rect._element.addprevious(pic._element)
    _remove(rect)


# ----------------------------------------------------------------- 표지
def fill_cover(slide, yy, mm, created, period):
    sm, sd, em, ed = period
    done = set()
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                t = r.text
                if "ST MM" in t:
                    r.text = t.replace("MM", f"{em:02d}"); done.add("title")
                elif t == "YY":
                    r.text = yy; done.add("yy")
                elif t.startswith("-MM-DD"):
                    r.text = t.replace("-MM-DD", f"-{created.month:02d}-{created.day:02d}"); done.add("date")
                elif t.startswith(": MM/DD"):
                    r.text = f": {sm:02d}/{sd:02d} - {em:02d}/{ed:02d}"; done.add("period")
    missing = {"title", "yy", "date", "period"} - done
    if missing:
        raise RuntimeError(f"표지에서 채우지 못한 칸: {sorted(missing)} — 양식 표지가 바뀌었는지 확인해 주세요.")


# ----------------------------------------------------------------- 콘텐츠 슬라이드
def _move(shape, dx):
    shape.left = int(shape.left + dx)


def _apply_split(slide, slots, n_hi, n_lo, group_labels):
    top = list(slide.shapes)
    bg = next(s for s in top if s.shape_type == 1 and not (s.has_text_frame and s.text_frame.text.strip()) and s.width > Inches(5))
    bg2_el = copy.deepcopy(bg._element)
    bg._element.addnext(bg2_el)
    bg2 = next(s for s in slide.shapes if s._element is bg2_el)
    gap = int(Inches(0.03))
    L0, W = bg.left, bg.width
    wl = int((W - gap) * n_hi / (n_hi + n_lo))
    regions = [(L0, wl, slots[:n_hi]), (L0 + wl + gap, W - gap - wl, slots[n_hi:])]
    for (rl, rw, group), bgs, text in zip(regions, (bg, bg2), group_labels):
        bgs.left, bgs.width = rl, rw
        cell = rw / max(len(group), 1)
        for k, sl in enumerate(group):
            dx = rl + cell * (k + 0.5) - _cx(sl["rect"])
            for key in ("rect", "label", "group"):
                _move(sl[key], dx)
        # 그룹 라벨: 영역의 첫 슬롯 라벨을 영역 전체 폭으로 늘려 쓰고 나머지는 지운다
        first = group[0]["label"]
        set_lines(first, [text])
        first.left, first.width = int(rl + Inches(0.3)), int(rw - Inches(0.6))
        for sl in group[1:]:
            _remove(sl["label"])
            sl["label"] = None


def fill_content(prs, cc, kind, cs, imgs, tpl_slides, log, note_text=None):
    n = max(3, min(5, len(cs.slots)))
    src = tpl_slides[SLOT_TEMPLATE_INDEX[n]]
    slide = new_slide(prs, f"{cc} - {kind}")
    clone_shapes(src, slide)
    for sh in list(slide.shapes):
        if sh.has_text_frame:
            t = sh.text_frame.text
            if "자리가 없을 경우" in t:
                _remove(sh)
            elif t.startswith("(국가)"):
                set_lines(sh, [cs.title])
            elif t.startswith("서브헤더는"):
                clear_text(sh)   # 코멘트 박스는 비워 둔다
    slots = find_slots(slide)
    for sl in slots[len(cs.slots):]:   # 남는 슬롯 삭제
        for key in ("rect", "label", "group"):
            _remove(sl[key])
    slots = slots[:len(cs.slots)]
    for sl, sp in zip(slots, cs.slots):
        set_lines(sl["label"], [sp.label])
        set_lines(sl["kpi"], sp.kpi if len(sl["kpi"].text_frame.paragraphs) >= 2 else [" / ".join(sp.kpi)])
        set_lines(sl["name"], split_sku(sp.name))
        img = imgs.find(sp.name)
        if img:
            put_image(slide, sl["rect"], img)
        else:
            log.append(f"이미지 없음: {sp.name}  ({cs.title.strip()})")
    if cs.split:
        _apply_split(slide, slots, cs.n_hi, cs.n_lo, cs.group_labels)
    if note_text:
        notes.write(slide, note_text)
    return slide


# ----------------------------------------------------------------- 표 슬라이드
def _set_cell_fill(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    for el in list(tcPr):
        if etree.QName(el).localname in ("solidFill", "noFill", "gradFill", "pattFill", "blipFill", "grpFill"):
            tcPr.remove(el)
    if not hexcolor:
        return
    sf = etree.Element(f"{{{A_NS}}}solidFill")
    etree.SubElement(sf, f"{{{A_NS}}}srgbClr").set("val", hexcolor)
    last = None
    for el in tcPr:
        if etree.QName(el).localname.startswith("ln"):
            last = el
    if last is not None:
        last.addnext(sf)
    else:
        tcPr.insert(0, sf)


def _fill_table(gf, spec):
    tbl = gf.table
    ncols = len(tbl.columns)
    if len(spec.headers) + 1 != ncols:
        raise ValueError(f"표 열 수가 맞지 않습니다: 시트 {len(spec.headers) + 1}열 / 양식 {ncols}열 ({spec.label})")
    need = len(spec.rows) + 1
    trs = list(tbl._tbl.tr_lst)
    while len(trs) > need:
        tbl._tbl.remove(trs.pop())
    while len(trs) < need:
        new = copy.deepcopy(trs[-1])
        trs[-1].addnext(new)
        trs.append(new)
    header = ["소재명"] + spec.headers
    for c, t in enumerate(header):
        set_para_text(tbl.cell(0, c).text_frame.paragraphs[0], t)
    for r, (name, vals, cls) in enumerate(spec.rows, start=1):
        for c, t in enumerate([name] + vals):
            cell = tbl.cell(r, c)
            set_para_text(cell.text_frame.paragraphs[0], t)
            _set_cell_fill(cell, GREEN if cls == "hi" else PINK if cls == "lo" else None)
    gf.height = Emu(sum(r.height for r in tbl.rows))


def fill_table_slide(prs, cc, kind, specs, tpl_slides):
    slide = new_slide(prs, f"{cc} - {kind}")
    clone_shapes(tpl_slides[TABLE_TEMPLATE_INDEX], slide)
    zone = "상시" if kind == "OG" else "프로모션"
    heads = sorted([s for s in slide.shapes if s.has_text_frame and s.text_frame.text.startswith("▼")], key=lambda s: s.top)
    tables = sorted([s for s in slide.shapes if s.has_table], key=lambda s: s.top)
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("SKIN1004"):
            set_lines(sh, [f"SKIN1004 {cc} {zone} 캠페인별 우수 소재 상세 효율"])
    for i, (h, t) in enumerate(zip(heads, tables)):
        if i < len(specs):
            n = len(specs[i].rows)
            set_lines(h, [f"▼ <{specs[i].label}> 소재 지출 금액 TOP {n}"])
            _fill_table(t, specs[i])
        else:
            _remove(h)
            _remove(t)
    return slide


# ----------------------------------------------------------------- 한 벌 만들기
def _find_by_text(slides, layout, text):
    for s in slides:
        if s.slide_layout.name == layout and any(sh.has_text_frame and sh.text_frame.text.strip() == text for sh in s.shapes):
            return s
    raise KeyError(f"양식에서 '{layout}' 레이아웃의 '{text}' 슬라이드를 찾지 못했습니다. 지원 국가: ID, MY, PH, SG")


def build_deck(plan, template_path, meta, images_dir, out_path, log, prev_sources=None):
    prs = Presentation(str(template_path))
    tpl = list(prs.slides)
    cc, kind = plan.country, plan.kind
    imgs = ImageIndex(images_dir)

    cover = tpl[0]
    fill_cover(cover, meta["yy"], meta["mm"], meta["created"], meta["period"])
    country_slide = _find_by_text(tpl, f"{cc} - OG", cc)
    kind_slide = _find_by_text(tpl, f"{cc} - {kind}", kind)

    order = [cover, country_slide, kind_slide]
    for cs in plan.slides:
        text = notes.render(plan, cs, meta, prev_sources) if prev_sources is not None else None
        order.append(fill_content(prs, cc, kind, cs, imgs, tpl, log, text))
    for i in range(0, len(plan.tables), 2):
        order.append(fill_table_slide(prs, cc, kind, plan.tables[i:i + 2], tpl))
    eod = new_slide(prs, f"{cc} - {kind}")
    clone_shapes(tpl[EOD_INDEX], eod)
    order.append(eod)

    keep = {id(s.part) for s in order}
    for s in tpl:
        if id(s.part) not in keep:
            delete_slide(prs, s)
    move_slide_to_end_order(prs, order)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path
