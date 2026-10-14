"""python-pptx 보조 함수: 슬라이드 복제·삭제, run 단위 텍스트 교체, 그룹 안 도형 찾기."""
import copy
import re

from pptx.util import Emu

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def layout_by_name(prs, name):
    for lay in prs.slide_layouts:
        if lay.name == name:
            return lay
    raise KeyError(f"레이아웃 '{name}' 이(가) 양식에 없습니다. 양식 pptx 를 확인해 주세요.")


def new_slide(prs, layout_name):
    """지정한 레이아웃의 빈 슬라이드(플레이스홀더 제거)."""
    s = prs.slides.add_slide(layout_by_name(prs, layout_name))
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def clone_shapes(src, dst):
    """src 슬라이드의 모든 도형을 dst 슬라이드로 깊은 복사한다 (그림 관계가 있으면 거부)."""
    for sh in src.shapes:
        xml = sh._element
        if xml.xpath('.//*[@r:embed or @r:link]'):
            raise ValueError("그림이 들어 있는 슬라이드는 복제할 수 없습니다 (양식에는 그림이 없어야 합니다).")
        dst.shapes._spTree.append(copy.deepcopy(xml))


def delete_slide(prs, slide):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        if prs.part.related_part(sldId.rId) is slide.part:
            prs.part.drop_rel(sldId.rId)
            sldIdLst.remove(sldId)
            return
    raise ValueError("삭제할 슬라이드를 찾지 못했습니다.")


def move_slide_to_end_order(prs, slides_in_order):
    """slides_in_order 에 준 순서대로 슬라이드를 재배열한다."""
    lst = prs.slides._sldIdLst
    by_part = {}
    for sldId in list(lst):
        by_part[id(prs.part.related_part(sldId.rId))] = sldId
    for s in slides_in_order:
        el = by_part[id(s.part)]
        lst.remove(el)
        lst.append(el)


def set_para_text(para, text):
    """문단의 첫 run 서식을 살려 텍스트를 바꾸고 나머지 run 은 지운다 (text_frame.text= 금지 규칙)."""
    runs = list(para.runs)
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        return
    r = para.add_run()
    r.text = text
    end = para._p.find(f"{{{A_NS}}}endParaRPr")
    if end is not None:
        rpr = copy.deepcopy(end)
        rpr.tag = f"{{{A_NS}}}rPr"
        old = r._r.find(f"{{{A_NS}}}rPr")
        if old is not None:
            r._r.remove(old)
        r._r.insert(0, rpr)


def set_lines(shape, lines):
    """shape 의 문단 개수를 lines 에 맞추고 각 문단 텍스트를 run 단위로 교체."""
    tf = shape.text_frame
    paras = list(tf.paragraphs)
    # 모자라면 마지막 문단을 복사해서 늘린다
    while len(paras) < len(lines):
        newp = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(newp)
        paras = list(tf.paragraphs)
    # 남으면 지운다
    for p in paras[len(lines):]:
        p._p.getparent().remove(p._p)
    for p, t in zip(list(tf.paragraphs), lines):
        set_para_text(p, t)


def clear_text(shape):
    """코멘트 박스처럼 비워 둘 도형: 문단 수·서식은 두고 글자만 지운다."""
    for p in shape.text_frame.paragraphs:
        set_para_text(p, "")


def all_shapes(shapes):
    """그룹 안까지 평탄화."""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from all_shapes(sh.shapes)


def text_of(sh):
    return sh.text_frame.text if sh.has_text_frame else ""


def split_sku(name):
    """소재명을 SKU 토큰(5번째 토큰) 뒤에서 2줄로 끊는다. 토큰이 5개 이하면 한 줄."""
    toks = name.split("_")
    if len(toks) <= 5:
        return [name]
    first = "_".join(toks[:5]) + "_"
    rest = "_".join(toks[5:])
    return [first, rest]
