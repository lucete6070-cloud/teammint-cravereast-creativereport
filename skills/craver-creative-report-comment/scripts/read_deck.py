"""생성기가 만든 pptx 에서 코멘트 작성 대상 슬라이드를 읽는다.

사용: python read_deck.py <deck.pptx> <출력폴더>
출력: <출력폴더>/deck.json, <출력폴더>/images/s{슬라이드}_slot{슬롯}.{확장자}
- 대상 = 노트가 '=== 코멘트 근거' 로 시작하는 슬라이드 (생성기 v1.2.0 이상)
- 슬롯 이미지는 소재명 글상자와 가로 위치가 가장 가까운 그림으로 짝짓는다 (없으면 null = 회색 사각형)
"""
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HEADER = "=== 코멘트 근거"
SLOT_RE = re.compile(r"^\[슬롯 (\d+)\] (\S+) \| (.*?) \| (\S+)$")


def flat(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from flat(sh.shapes)


def comment_box(slide):
    for sh in slide.shapes:
        if sh.name == "코멘트 박스" and sh.has_text_frame:
            return sh
    return None


def parse_notes(text):
    info = {"slots": []}
    cur = None
    for line in text.splitlines():
        if line.startswith("리포트: "):
            info["report"] = line[5:]
        elif line.startswith("슬라이드: "):
            info["slide"] = line[6:]
        elif line.startswith("제목: "):
            info["title"] = line[4:]
        elif line.startswith("전월 비교 자료: "):
            info["prev_sources"] = line[10:]
        m = SLOT_RE.match(line)
        if m:
            cur = {"no": int(m.group(1)), "grade": m.group(2), "label": m.group(3), "name": m.group(4)}
            info["slots"].append(cur)
        elif cur and line.startswith("  지표: "):
            cur["metrics"] = line[6:]
        elif cur and line.startswith("  KPI: "):
            cur["metrics"] = line[7:]
        elif cur and line.startswith("  시트 메모: "):
            cur["memo"] = line[9:]
        elif cur and line.startswith("  전월: "):
            cur["prev"] = line[6:]
        elif line.startswith("=== 코멘트 스킬 메모"):
            cur = None
    return info


def main(deck, out):
    out = Path(out)
    (out / "images").mkdir(parents=True, exist_ok=True)
    prs = Presentation(deck)
    result = {"deck": str(deck), "slides": []}
    for idx, slide in enumerate(prs.slides, 1):
        if not slide.has_notes_slide:
            continue
        notes = slide.notes_slide.notes_text_frame.text
        if not notes.startswith(HEADER):
            continue
        info = parse_notes(notes)
        box = comment_box(slide)
        info.update(index=idx, notes=notes, comment_box=bool(box),
                    comment_box_text=box.text_frame.text.strip() if box else None)
        shapes = list(flat(slide.shapes))
        pics = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        for slot in info["slots"]:
            key = slot["name"].replace("\n", "")
            name_shape = next((s for s in shapes if s.has_text_frame
                               and s.text_frame.text.replace("\n", "").replace("\v", "").strip() == key), None)
            slot["image"] = None
            if name_shape is not None and pics:
                cx = name_shape.left + name_shape.width / 2
                pic = min(pics, key=lambda p: abs(p.left + p.width / 2 - cx))
                if abs(pic.left + pic.width / 2 - cx) < pic.width:
                    f = out / "images" / f"s{idx}_slot{slot['no']}.{pic.image.ext}"
                    f.write_bytes(pic.image.blob)
                    slot["image"] = str(f)
        result["slides"].append(info)
    (out / "deck.json").write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
    n_img = sum(1 for s in result["slides"] for sl in s["slots"] if sl["image"])
    n_slot = sum(len(s["slots"]) for s in result["slides"])
    print(f"대상 슬라이드 {len(result['slides'])}장, 슬롯 {n_slot}개 (이미지 {n_img}개) → {out / 'deck.json'}")
    if not result["slides"]:
        print("[멈춤] 노트에 '코멘트 근거'가 있는 슬라이드가 없습니다. 생성기 v1.2.0 이상으로 만든 pptx 인지 확인하세요.")
        return 1
    if any(not s["comment_box"] for s in result["slides"]):
        print("[주의] 일부 슬라이드에서 '코멘트 박스'를 찾지 못했습니다. 생성기 v1.2.1 이상으로 다시 만들어 주세요.")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1], sys.argv[2]))
