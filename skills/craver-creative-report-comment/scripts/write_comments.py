"""코멘트 초안(drafts.json)을 pptx 코멘트 박스와 노트에 넣는다. 넣기 전에 형식 점검을 한다.

사용: python write_comments.py <deck.pptx> <drafts.json> <출력.pptx> [--check-only]

drafts.json 형식:
{"slides": [{"index": 4,                      # read_deck.py 의 index (1부터)
             "lines": ["본문1", "본문2", "본문3"], # 2~4줄
             "arrow": "▶ 제언",                # 생략 가능 (생략 시 플래그)
             "flags": [{"line": 2, "type": "근거 약함", "reason": "..."}],
             "callouts": [{"slot": 2, "text": "설명 요소 비중 높음", "reason": "..."}],
             "basis": ["P-MY-03", "판정 MY J-01"]}]}

- 코멘트 박스 서식 보존: 본문 줄은 첫 문단, ▶ 줄은 마지막 문단(금색) 서식을 복사해 run 단위로 쓴다.
- 노트의 '=== 코멘트 스킬 메모' 아래에 근거·플래그·콜아웃 후보·자동 점검 결과를 적는다.
- 숫자 점검: 본문에 쓴 %·₩·원 숫자가 노트(시트 값)에 없으면 플래그.
"""
import copy
import json
import re
import sys
from pathlib import Path

from pptx import Presentation

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
SKILL = "=== 코멘트 스킬 메모"
ASSERTIVE = ("때문", "원인은", "확실", "분명히", "반드시")
NUM = re.compile(r"[₩$]?\d[\d,.]*\s*(?:%p|%|원|만원)?")


def box_of(slide):
    for sh in slide.shapes:
        if sh.name == "코멘트 박스" and sh.has_text_frame:
            return sh
    return None


def lint(d, notes):
    issues = []
    lines = d.get("lines") or []
    if not 2 <= len(lines) <= 4:
        issues.append(f"본문 {len(lines)}줄 (2~4줄이어야 함)")
    arrow = (d.get("arrow") or "").strip()
    if not arrow:
        issues.append("▶ 줄 없음")
    elif not arrow.startswith("▶"):
        issues.append("▶ 줄이 ▶ 로 시작하지 않음")
    for i, t in enumerate(lines + ([arrow] if arrow else []), 1):
        tag = "▶" if arrow and i == len(lines) + 1 else f"{i}줄"
        if len(t) > 115:
            issues.append(f"{tag}: {len(t)}자 (115자 이하 권장)")
        if t.rstrip().endswith(("다.", "다", "요.", ",")):
            issues.append(f"{tag}: 개조식 명사형 종결이 아님 ('{t.rstrip()[-4:]}')")
        for w in ASSERTIVE:
            if w in t:
                issues.append(f"{tag}: 단정 표현 '{w}' (R1)")
        for n in NUM.findall(t):
            n = n.strip()
            if re.search(r"[%₩원]", n) and n.replace(" ", "") not in notes.replace(" ", ""):
                issues.append(f"{tag}: 숫자 '{n}' 가 시트 값(노트)에 없음 — 확인 필요")
    ends = [re.sub(r"[\s.]+$", "", t)[-2:] for t in lines]
    for e in set(ends):
        if ends.count(e) >= 3:
            issues.append(f"같은 어미 '{e}' {ends.count(e)}번 반복")
    return issues


def _para_with_text(tpl_p, text):
    p = copy.deepcopy(tpl_p)
    runs = p.findall(f"{A}r")
    if not runs:
        raise ValueError("코멘트 박스 문단에 서식 run 이 없습니다.")
    runs[0].find(f"{A}t").text = text
    for r in runs[1:]:
        p.remove(r)
    return p


def fill_box(box, lines, arrow):
    txBody = box.text_frame._txBody
    paras = txBody.findall(f"{A}p")
    body_tpl, arrow_tpl = paras[0], paras[-1]
    if not body_tpl.findall(f"{A}r") or not arrow_tpl.findall(f"{A}r"):
        raise ValueError("코멘트 박스 서식이 예상과 다릅니다 (빈 문단). 생성기로 다시 만든 pptx 를 쓰세요.")
    new = [_para_with_text(body_tpl, t) for t in lines]
    if arrow:
        new.append(_para_with_text(arrow_tpl, arrow))
    for p in paras:
        txBody.remove(p)
    for p in new:
        txBody.append(p)


def memo(d, issues):
    out = []
    if d.get("basis"):
        out.append("근거: " + ", ".join(d["basis"]))
    for f in d.get("flags", []):
        out.append(f"[{f.get('type', '플래그')}] {f.get('line', '')}줄: {f.get('reason', '')}")
    for c in d.get("callouts", []):
        out.append(f"[콜아웃 후보] 슬롯 {c.get('slot')}: {c.get('text')} — {c.get('reason', '')}")
    for i in issues:
        out.append(f"[자동 점검] {i}")
    return out


def main(deck, drafts, out, check_only=False):
    prs = Presentation(deck)
    slides = list(prs.slides)
    data = json.loads(Path(drafts).read_text(encoding="utf-8"))
    report, bad = [], 0
    for d in data["slides"]:
        s = slides[d["index"] - 1]
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        if not notes.startswith("=== 코멘트 근거"):
            print(f"[멈춤] {d['index']}번 슬라이드는 코멘트 대상(노트에 코멘트 근거)이 아닙니다.")
            return 1
        box = box_of(s)
        if box is None:
            print(f"[멈춤] {d['index']}번 슬라이드에서 '코멘트 박스'를 찾지 못했습니다.")
            return 1
        issues = lint(d, notes.split(SKILL)[0])   # 숫자 대조는 생성기가 적은 시트 값 부분만
        bad += bool(issues)
        report.append((d["index"], issues, d.get("flags", [])))
        if check_only:
            continue
        fill_box(box, d["lines"], (d.get("arrow") or "").strip())
        body = notes.split(SKILL)[0].rstrip("\n")
        tail = "\n".join(memo(d, issues))
        s.notes_slide.notes_text_frame.text = f"{body}\n\n{SKILL} (초안 근거·플래그·콜아웃 후보) ===\n{tail}\n"
    for idx, issues, flags in report:
        print(f"- {idx}번 슬라이드: 자동 점검 {len(issues)}건, 스킬 플래그 {len(flags)}건")
        for i in issues:
            print(f"    · {i}")
    if not check_only:
        Path(out).parent.mkdir(parents=True, exist_ok=True)
        prs.save(out)
        print(f"저장: {out}")
    return 0


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--check-only"]
    if len(args) != 3:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(*args, check_only="--check-only" in sys.argv))
