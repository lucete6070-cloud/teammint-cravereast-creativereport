"""코멘트 스킬 스크립트 테스트 (합성 데이터). 실행: python tests/test_skill_scripts.py

생성기 → read_deck.py → (초안 JSON) → write_comments.py 순서로 돌려
코멘트 박스에 초안이 들어가고 ▶ 줄 금색 서식이 유지되는지, 노트에 메모·자동 점검이 붙는지 본다.
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))
import run  # noqa: E402
from synth import A1, A2, A3, P1, P2, P3, make_inputs  # noqa: E402

SK = ROOT / "skills" / "craver-creative-report-comment" / "scripts"


def py(*args):
    r = subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True, encoding="utf-8")
    print(r.stdout[-1500:], r.stderr[-1500:])
    return r


with tempfile.TemporaryDirectory() as d:
    d = Path(d)
    make_inputs(d)
    (d / "images").mkdir()
    for i, n in enumerate([A1, A2, A3, P1, P2, P3]):
        Image.new("RGB", (400, 400), (40 * i, 100, 200)).save(d / "images" / f"{n}.png")
    assert run.main(["--input", str(d / "input"), "--images", str(d / "images"), "--out", str(d / "out"),
                     "--history", str(d / "ref"), "--created", "2026-10-25"]) == 0
    deck = d / "out" / "SKIN1004_26-10_EAST_MY_OG.pptx"

    r = py(SK / "read_deck.py", deck, d / "work")
    assert r.returncode == 0
    info = json.loads((d / "work" / "deck.json").read_text(encoding="utf-8"))
    assert len(info["slides"]) == 3 and all(s["comment_box"] for s in info["slides"])
    atc = next(s for s in info["slides"] if "유입" in s["slide"])
    assert [sl["name"] for sl in atc["slots"]] == [A1, A2, A3]
    assert atc["slots"][1]["prev"].startswith("26-09 저조 (전월 시트)")
    # 슬롯 이미지가 그 소재의 그림과 짝지어졌는지 (색으로 확인)
    for k, sl in enumerate(atc["slots"]):
        assert sl["image"], sl
        assert Image.open(sl["image"]).convert("RGB").getpixel((5, 5))[0] == 40 * [A1, A2, A3].index(sl["name"])

    drafts = {"slides": [{
        "index": atc["index"],
        "lines": ["우수 소재는 할인율과 판매가를 한눈에 전면 배치해 구매 메리트를 직관적으로 전달한 양상",
                  "TNB 앰플 소재는 이번 달 포함 3개월 연속 저조군에 포함되며 설명 요소 중심 구성이 유지된 경향성",
                  "저조 소재는 우수 소재 대비 효능 설명 비중이 높아 가격 주목도가 낮았을 가능성 (ROAS 312%)"],
        "arrow": "▶ 첫 화면은 가격·구성 중심으로 단순화하고 반복 저조 소재는 교체 검토 제언",
        "flags": [{"line": 3, "type": "근거 약함", "reason": "이미지상 효능 문구 비중 추정"}],
        "callouts": [{"slot": 2, "text": "설명 요소 비중 높음", "reason": "효능 문구가 첫 화면 절반"}],
        "basis": ["P-MY-03", "판정 MY J-01"]}, {
        "index": next(s["index"] for s in info["slides"] if "구매 / 저조" in s["slide"]),
        "lines": ["신제품 소재는 구매 이유 제시가 부족했기 때문에 저조한 것이다.", "ROAS 999% 기록"],
        "arrow": ""}]}
    (d / "work" / "drafts.json").write_text(json.dumps(drafts, ensure_ascii=False), encoding="utf-8")

    r = py(SK / "write_comments.py", deck, d / "work" / "drafts.json", d / "out" / "draft.pptx", "--check-only")
    assert r.returncode == 0 and not (d / "out" / "draft.pptx").exists()
    assert "단정 표현 '때문'" in r.stdout and "▶ 줄 없음" in r.stdout and "'999%'" in r.stdout and "개조식" in r.stdout
    r = py(SK / "write_comments.py", deck, d / "work" / "drafts.json", d / "out" / "draft.pptx")
    assert r.returncode == 0

    prs = Presentation(str(d / "out" / "draft.pptx"))
    s = list(prs.slides)[atc["index"] - 1]
    box = next(sh for sh in s.shapes if sh.name == "코멘트 박스")
    paras = box.text_frame.paragraphs
    assert [p.text for p in paras] == drafts["slides"][0]["lines"] + [drafts["slides"][0]["arrow"]]
    assert paras[-1].runs[0].font.color.rgb == (0x90, 0x73, 0x44)          # ▶ 줄 금색 유지
    assert paras[0].runs[0].font.color.type is None                       # 본문은 기본색
    n = s.notes_slide.notes_text_frame.text
    assert n.startswith("=== 코멘트 근거") and "[콜아웃 후보] 슬롯 2: 설명 요소 비중 높음" in n
    assert "근거: P-MY-03, 판정 MY J-01" in n and "[근거 약함] 3줄" in n
    assert n.count("=== 코멘트 스킬 메모") == 1
    # 다시 넣어도 메모가 겹치지 않는다
    r = py(SK / "write_comments.py", d / "out" / "draft.pptx", d / "work" / "drafts.json", d / "out" / "draft2.pptx")
    assert "'999%'" in r.stdout, "다시 넣을 때도 시트에 없는 숫자를 잡아야 함"
    n2 = list(Presentation(str(d / "out" / "draft2.pptx")).slides)[atc["index"] - 1].notes_slide.notes_text_frame.text
    assert n2.count("=== 코멘트 스킬 메모") == 1 and n2.count("[콜아웃 후보]") == 1
print("OK")
