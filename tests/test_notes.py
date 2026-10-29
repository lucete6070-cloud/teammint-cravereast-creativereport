"""코멘트 근거 노트 테스트 (합성 데이터, 광고주 데이터 없음). 실행: python tests/test_notes.py

이번 달 MY OG 시트 + 전월 시트 + 소재이력 파일로 생성기를 돌리고, 콘텐츠 슬라이드 노트에
슬롯·시트 지표(문자열 그대로)·전월 기록·2개월 연속 표시가 들어가는지, 코멘트 박스는 비어 있는지 본다.
"""
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

import run  # noqa: E402
from pptx import Presentation  # noqa: E402
sys.path.insert(0, str(ROOT / "tests"))
from synth import A1, A2, make_inputs  # noqa: E402,F401

with tempfile.TemporaryDirectory() as d:
    d = Path(d)
    make_inputs(d)
    rc = run.main(["--input", str(d / "input"), "--images", str(d / "images"), "--out", str(d / "out"),
                   "--history", str(d / "ref"), "--created", "2026-10-25"])
    assert rc == 0
    log = (d / "out" / "생성로그.txt").read_text(encoding="utf-8")
    print(log)
    assert "전월 비교 자료 = 전월 시트 CSV + 소재이력_MY.md" in log
    prs = Presentation(str(d / "out" / "SKIN1004_26-10_EAST_MY_OG.pptx"))
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    notes = [n for n in notes if n.startswith("=== 코멘트 근거")]
    assert len(notes) == 3, len(notes)   # 유입 통합 1 + 구매 우수 1 + 구매 저조 1
    atc = next(n for n in notes if "유입" in n.splitlines()[2])
    print(atc)
    assert "MY 상시(OG) 유입 / 우수<>저조" in atc
    assert "지출금액 ₩910,734" in atc and "ROAS 1225%" in atc          # 시트 문자열 그대로
    assert f"[슬롯 2] 저조 | ATC | {A2}" in atc
    assert "26-09 저조 (전월 시트) → 이번 달 포함 3개월 연속 저조 (소재이력 기준)" in atc   # 10월+9월(전월 시트)+8월(이력)
    assert "26-09 우수 (전월 시트) → 이번 달 포함 2개월 연속 우수" in atc
    assert "26-09 집행(체크 없음) (전월 시트)" in atc                       # A3: 전월 체크 없음
    pur_lo = next(n for n in notes if "구매 / 저조" in n)
    assert "26-09 저조 (전월 시트) → 이번 달 포함 2개월 연속 저조" in pur_lo
    pur_hi = next(n for n in notes if "구매 / 우수" in n)
    assert "26-09 전월 시트에 없음" in pur_hi                                 # P3 신규
    assert "=== 코멘트 스킬 메모" in atc
    # 코멘트 박스(서브헤더 자리)는 여전히 비어 있다
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                assert not sh.text_frame.text.startswith("서브헤더는")

    # 소재이력만 있을 때 (보조가 전월 시트를 안 넣은 경우)
    rc = run.main(["--input", str(d / "input"), "--images", str(d / "images"), "--out", str(d / "out3"),
                   "--prev", str(d / "none"), "--history", str(d / "ref"), "--created", "2026-10-25"])
    assert rc == 0
    prs = Presentation(str(d / "out3" / "SKIN1004_26-10_EAST_MY_OG.pptx"))
    t = "\n".join(s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide)
    assert "전월 비교 자료: 소재이력_MY.md" in t
    assert "26-09 저조 (소재이력) → 이번 달 포함 3개월 연속 저조 (소재이력 기준)" in t
    assert "26-09 기록 없음 (소재이력 기준)" in t            # 이력에 없는 A1

    # 전월 자료가 하나도 없으면 '없음'으로 적고 생성은 계속
    rc = run.main(["--input", str(d / "input"), "--images", str(d / "images"), "--out", str(d / "out2"),
                   "--prev", str(d / "none"), "--history", str(d / "none"), "--created", "2026-10-25"])
    assert rc == 0
    prs = Presentation(str(d / "out2" / "SKIN1004_26-10_EAST_MY_OG.pptx"))
    n0 = next(t for t in (s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide) if t.startswith("=== 코멘트 근거"))
    assert "전월 비교 자료: 없음" in n0 and "전월 자료 없음" in n0
print("OK")
